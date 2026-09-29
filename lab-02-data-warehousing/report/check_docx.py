#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Автоматическая проверка структуры отчета DOCX.

Проверяется:
  * формат страниц, поля, ориентация разделов;
  * стили заголовков и наличие оглавления;
  * нумерация страниц в колонтитулах (титул без номера);
  * ширина колонок таблиц относительно самого длинного слова в заголовке
    (защита от переносов вида "Нормализац / ия");
  * наличие рисунков и подписей к ним;
  * отсутствие placeholder-текста и незаполненных мест.

Запуск:  python3 report/check_docx.py
"""
import os
import re
import sys
import zipfile
from collections import Counter

from PIL import ImageFont
from docx import Document
from docx.shared import Cm

ZWSP = "\u200b"

HERE = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.join(os.path.dirname(HERE), "Отчет", "Лабораторная работа №2.docx")
F_BOLD = "/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf"
F_REG = "/System/Library/Fonts/Supplemental/Times New Roman.ttf"

CELL_PAD_CM = 45 / 567 * 2 + 0.05   # отступы ячейки, заданные в docx_helpers

problems = []
notes = []


def word_cm(word, size_pt, bold):
    """Ширина строки в сантиметрах для шрифта Times New Roman."""
    f = ImageFont.truetype(F_BOLD if bold else F_REG, int(size_pt * 4))
    return f.getlength(word) / 4 / 72 * 2.54


def check_sections(doc):
    print("1. Разделы документа")
    for i, s in enumerate(doc.sections):
        w = s.page_width.cm
        h = s.page_height.cm
        orient = "landscape" if w > h else "portrait"
        print(f"   раздел {i}: {w:.1f} x {h:.1f} см ({orient}), поля "
              f"L{s.left_margin.cm:.1f} R{s.right_margin.cm:.1f} "
              f"T{s.top_margin.cm:.1f} B{s.bottom_margin.cm:.1f}")
        if orient == "portrait":
            expect = (3.0, 1.0, 2.0, 2.0)
            got = (round(s.left_margin.cm, 1), round(s.right_margin.cm, 1),
                   round(s.top_margin.cm, 1), round(s.bottom_margin.cm, 1))
            if got != expect:
                problems.append(f"раздел {i}: поля {got} вместо {expect}")
    if len(doc.sections) < 2:
        problems.append("нет отдельного раздела для альбомных приложений")


def check_footers(doc):
    print("2. Нумерация страниц")
    for i, s in enumerate(doc.sections):
        xml = s.footer._element.xml if s.footer is not None else ""
        has_page = "PAGE" in xml
        print(f"   раздел {i}: номер страницы {'есть' if has_page else 'нет'}")
        if i == 0 and has_page:
            problems.append("титульный раздел содержит номер страницы")
        if i > 0 and not has_page:
            problems.append(f"раздел {i} без номера страницы")


def check_toc(doc):
    print("3. Оглавление")
    xml = doc.element.xml
    if "TOC \\o" not in xml and 'TOC \\o' not in xml:
        problems.append("не найдено поле оглавления TOC")
    else:
        print("   поле TOC присутствует")
    n_cached = len(re.findall(r"<w:tab/>", xml))
    print("   кэшированных строк оглавления:", n_cached)
    if n_cached < 8:
        problems.append("оглавление не содержит кэшированных строк")


def check_headings(doc):
    print("4. Иерархия заголовков")
    c = Counter(p.style.name for p in doc.paragraphs if p.style.name.startswith("Heading"))
    print("  ", dict(c))
    if c.get("Heading 1", 0) < 8:
        problems.append("слишком мало заголовков уровня 1")
    if c.get("Heading 2", 0) < 5:
        problems.append("слишком мало заголовков уровня 2")
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading") and not p.text.strip():
            problems.append("пустой заголовок")


def check_tables(doc):
    print("5. Ширина колонок таблиц")
    bad = 0
    for ti, t in enumerate(doc.tables):
        if len(t.rows) == 0:
            continue
        widths = []
        for i, col in enumerate(t.columns):
            widths.append(col.width.cm if col.width else None)
        hdr = [c.text for c in t.rows[0].cells]
        size = 12.0
        for c in t.rows[0].cells:
            for p in c.paragraphs:
                for r in p.runs:
                    if r.font.size:
                        size = r.font.size.pt
                        break
        for i, text in enumerate(hdr):
            if widths[i] is None:
                continue
            avail = widths[i] - CELL_PAD_CM
            longest = max(text.replace(ZWSP, " ").split(),
                          key=lambda w: word_cm(w, size, True)) if text.split() else ""
            need = word_cm(longest, size, True)
            if need > avail + 0.05:      # допуск на округление ширин
                print(f"   таблица {ti}, колонка {i}: «{longest}» нужно "
                      f"{need:.2f} см, доступно {avail:.2f} см")
                bad += 1
    if bad:
        problems.append(f"{bad} колонок с переносом внутри слова")
    else:
        print("   переносов внутри слов в заголовках нет")

    small = []
    for ti, t in enumerate(doc.tables):
        sizes = set()
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        if run.font.size:
                            sizes.add(round(run.font.size.pt, 1))
        if sizes and min(sizes) < 12:
            small.append((ti, sorted(sizes)))
    print("   таблицы мельче 12 pt:", small if small else "нет")
    problems.extend(("таблица мельче 12 pt", x) for x in small)


def check_images(docx_path):
    print("6. Рисунки и подписи")
    with zipfile.ZipFile(docx_path) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
    print(f"   файлов изображений в пакете: {len(media)}")
    doc = Document(docx_path)
    xml = doc.element.xml
    n_draw = xml.count("<w:drawing>")
    print(f"   вставок изображений в документе: {n_draw}")
    caps = re.findall(r"Рисунок [А-ЯA-Z]?\.?\d+[^<]{0,70}", xml)
    print(f"   подписей рисунков: {len(caps)}")
    for c in caps:
        print("     ", c[:80])
    if len(media) < 2:
        problems.append("в документе меньше двух уникальных изображений")
    if n_draw < 4:
        problems.append("в документе меньше четырех вставок изображений")
    if len(caps) < 4:
        problems.append("меньше четырех подписей к рисункам")


def check_placeholders(doc):
    print("7. Placeholders и незаполненные места")
    bad_words = ["TODO", "TBD", "XXX", "FIXME", "lorem", "Lorem",
                 "{{", "}}", "PLACEHOLDER", "placeholder", "<...>"]
    hits = []
    for p in doc.paragraphs:
        for w in bad_words:
            if w in p.text:
                hits.append((w, p.text[:70]))
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for w in bad_words:
                    if w in c.text:
                        hits.append((w, c.text[:70]))
    print("   найдено:", hits if hits else "нет")
    problems.extend(("placeholder", h) for h in hits)


def main():
    if not os.path.exists(DOCX):
        print("Нет файла:", DOCX)
        return 1
    doc = Document(DOCX)
    check_sections(doc)
    check_footers(doc)
    check_toc(doc)
    check_headings(doc)
    check_tables(doc)
    check_images(DOCX)
    check_placeholders(doc)

    print()
    print("=" * 70)
    if problems:
        print(f"ИТОГ: обнаружено {len(problems)} замечаний")
        for p in problems:
            print("  -", p)
        return 1
    print("ИТОГ: замечаний нет")
    return 0


if __name__ == "__main__":
    sys.exit(main())
